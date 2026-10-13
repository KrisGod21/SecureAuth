from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / 'docs' / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
FONT = 'C:/Windows/Fonts/calibri.ttf'
BOLD = 'C:/Windows/Fonts/calibrib.ttf'
INK = '#182B45'
BLUE = '#295FBA'

def flow(filename, nodes, decisions=None):
    decisions = decisions or {}
    W, H = (1660, 2050) if decisions else (1400, 1800)
    img = Image.new('RGB', (W,H), 'white')
    draw = ImageDraw.Draw(img)
    regular = ImageFont.truetype(FONT, 32)
    bold = ImageFont.truetype(BOLD, 33)
    small = ImageFont.truetype(FONT, 28)
    cx, left, right = (625, 165, 1085) if decisions else (700, 100, 1300)
    top, height = 30, 135
    gap = (H - 70 - len(nodes)*height) // (len(nodes)-1)
    def text_center(text, x, y, font=regular, color=INK):
        bbox=draw.multiline_textbbox((0,0),text,font=font,spacing=8,align='center')
        draw.multiline_text((x-(bbox[2]-bbox[0])/2,y-(bbox[3]-bbox[1])/2-4),text,font=font,fill=color,spacing=8,align='center')
    def arrow(x1,y1,x2,y2):
        draw.line((x1,y1,x2,y2),fill='#8796AB',width=4)
        if x1==x2: draw.polygon([(x2,y2),(x2-10,y2-16),(x2+10,y2-16)],fill='#8796AB')
        else: draw.polygon([(x2,y2),(x2-16,y2-10),(x2-16,y2+10)],fill='#8796AB')
    for i, text in enumerate(nodes):
        y=top+i*(height+gap)
        if i in decisions:
            draw.polygon([(cx,y),(right,y+height/2),(cx,y+height),(left,y+height/2)],fill='#F0F4FA',outline='#9BAFCB',width=3)
            text_center(text,cx,y+height/2,bold)
            arrow(right,y+height/2,1230,y+height/2)
            draw.rounded_rectangle((1230,y+15,1610,y+height-15),radius=22,fill='#FFF4F3',outline='#D1A39F',width=3)
            text_center('Reject\nNo plaintext returned',1420,y+height/2,small)
            text_center('No',1150,y+height/2-26,small)
            if i<len(nodes)-1: text_center('Yes',cx+50,y+height+gap/2,small)
        else:
            fill='#EAF1FD' if i in [0,len(nodes)-1] else '#F8FAFD'
            draw.rounded_rectangle((left,y,right,y+height),radius=35 if i in [0,len(nodes)-1] else 15,fill=fill,outline='#9BAFCB',width=3)
            text_center(text,cx,y+height/2,bold if i in [0,len(nodes)-1] else regular)
        if i<len(nodes)-1: arrow(cx,y+height,cx,y+height+gap)
    path=ASSETS/filename
    img.save(path)
    return path

encryption=flow('encryption.png',[
    'Input plaintext P, field context L\nand 256 bit master key K',
    'HKDF SHA256 derives\nKa, Kd, Ks and Km',
    'Compute synthetic IV V\nHMAC of framed context and plaintext',
    'Split V into two IVs\nIVa = V[0:16]   IVd = V[16:24]',
    'Encrypt P using AES 256 CTR\nA = AES CTR(Ka, IVa, P)',
    'Apply PKCS7 padding to A\nEncrypt using Triple DES CBC',
    'Build body = version || V || B\nMAC context and body with Km',
    'Encode body || MAC as Base64URL\nCompute SHA256 of encoded text',
    'Output packet\n{ ciphertext ct, digest }'
])
decryption=flow('decryption.png',[
    'Input ciphertext, digest, context and key',
    'Check digest and canonical encoding\nValidate version and lengths',
    'Derive keys and reconstruct body\nCompute expected outer MAC',
    'Outer MAC valid?',
    'Undo Triple DES CBC\nStrictly validate and remove PKCS7',
    'Undo AES 256 CTR\nRecover candidate plaintext P',
    'Synthetic IV valid?',
    'Decode UTF8 and return plaintext\nOnly after both checks pass'
],{3:True,6:True})
password_flow=flow('password.png',[
    'Optionally obtain 64 ANU quantum bytes\nValidate response and source status',
    'Generate 64 independent bytes\nfrom the operating system CSPRNG',
    'If ANU succeeds, XOR the two sources\nOtherwise retain CSPRNG bytes',
    'HKDF SHA256 expands the mixed seed\nUse password counter as context',
    'Accept byte b only when b < 210\nSelect alphabet[b mod 70]',
    'Build a 16 character candidate\nRequire upper, lower, digit and symbol',
    'If a class is missing, increment counter\nand derive a fresh candidate',
    'Hash password for storage with scrypt\nEncrypt the response for browser QR'
])

doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.27); section.page_height=Inches(11.69)
section.top_margin=Inches(.7); section.bottom_margin=Inches(.65)
section.left_margin=Inches(.8); section.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]; style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(27)
doc.styles['Title'].paragraph_format.space_after=Pt(13)
doc.styles['Subtitle'].font.size=Pt(12)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(11)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(6)
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
doc.core_properties.title='AC Algorithm and Application'
doc.core_properties.subject='HC ADS v2 encrypted registration and login'
doc.core_properties.author='AC Project'
doc.core_properties.keywords='cryptography, authenticated encryption, HC-ADS, registration, login'
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('AC Secure Auth Lab  |  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text, bold_start=None):
    para=doc.add_paragraph()
    if bold_start and text.startswith(bold_start):para.add_run(bold_start).bold=True;para.add_run(text[len(bold_start):])
    else:para.add_run(text)
    return para
def heading(text):doc.add_heading(text,level=2)
def page(title):
    para=doc.add_heading(title,level=1)
    para.paragraph_format.page_break_before=True
def formula(text):
    para=p(text);para.paragraph_format.space_after=Pt(7)
    for r in para.runs:r.font.name='Consolas';r.font.size=Pt(9)
    return para
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for cell,w in zip(t.rows[0].cells,widths):cell.width=Inches(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for cell,value,w in zip(cells,row,widths):cell.text=str(value);cell.width=Inches(w)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EDFA' if ri==0 else ('F6F8FC' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.05
                for run in para.runs:run.font.size=Pt(9);run.font.bold=ri==0
    spacer=doc.add_paragraph()
    spacer.paragraph_format.space_after=Pt(0)
    spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=Pt(2)
    spacer.add_run().font.size=Pt(1)
    return t
def image(path,caption,width=6.15):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=para.add_run();run.add_picture(str(path),width=Inches(width))
    pic=run._r.xpath('.//wp:docPr')[0];pic.set('descr',caption)
    para.paragraph_format.space_after=Pt(4)
    para=p(caption);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(9)

doc.add_paragraph('AC Algorithm and Application',style='Title')
doc.add_paragraph('Secure registration and login with HC ADS v2',style='Subtitle')
p('Experiment 9  |  Applied Cryptography  |  13 October 2026')
heading('Purpose and result')
p('The application demonstrates registration and login using encrypted fields, ciphertext fingerprints, generated passwords and a browser-rendered QR code. It includes a responsive interface, an Express API, SQLite storage and automated security tests. Profile fields are stored as authenticated ciphertext. The normalized student ID is also stored as a readable username; passwords are stored as salted scrypt hashes.')
p('HC-ADS v2 is a project-specific composition of AES-256-CTR, Triple-DES-CBC, HKDF-SHA256 and HMAC-SHA256. Its principal improvement over the original plan is an independent MAC verified before any decryption. A temporary browser channel key and a separate server-only storage key remove the need to expose one universal master key to clients.')
p('This is an educational cryptography implementation. The composition has no formal security proof or independent audit. Triple DES is retired for new NIST encryption use; the application must not be represented as production-certified security. Novelty here means the project-specific orchestration, not a new proven cryptographic primitive.')
heading('Application components')
table(['Component','Responsibility'],[
    ['Browser','Collect fields, negotiate a temporary key, encrypt and fingerprint requests, authenticate responses, display a local password QR and profile.'],
    ['Express server','Consume challenges, authenticate envelopes, validate plaintext in memory, generate passwords, encrypt profile fields, hash and verify passwords, manage sessions.'],
    ['SQLite database','Persist username, salted password hash, encrypted profile fields, a random row ID and a creation timestamp.'],
    ['Shared cipher module','One versioned implementation used by both the browser and Node to avoid divergent cipher formats.']
],[1.3,5.35])
heading('Changes from the supplied implementation plan')
p('The AES and Triple-DES cascade, SHA-256 comparison and QR delivery are preserved. The current password backend stores salted scrypt hashes and readable usernames. Version 2 adds a fourth derived key for an outer MAC, length-prefixed context encoding, temporary ECDH channel keys, one-use challenges, sessions and rate limits. Other personal fields use row-specific storage contexts. Native node:sqlite replaces a native addon, and CryptoJS supplies the portable legacy cipher layer. Genuine quantum randomness is optional; there is no simulated quantum claim.')

page('Keys and field contexts')
p('Two independent key domains separate messages from stored records. The storage key Kstore is a randomly generated 32-byte secret in the server .env file. The channel key Kchannel is a 32-byte secret established separately for each browser channel. Both are represented as exactly 64 hexadecimal characters when passed to the shared module.')
table(['Symbol','Size','Meaning'],[
    ['P','0 to 1024 bytes','UTF-8 plaintext bytes for one field.'],['L','1 to 512 bytes','UTF-8 field context authenticated with the envelope.'],['Ka','32 bytes','AES-256 key.'],['Kd','24 bytes','Triple-DES EDE3 key material; 192 stored bits, about 112-bit effective security.'],['Ks','32 bytes','Key for the synthetic-IV HMAC.'],['Km','32 bytes','Independent key for the outer authentication HMAC.'],['V','32 bytes','Synthetic-IV tag, derived from context and plaintext.'],['T','32 bytes','Outer authentication tag, covering context and ciphertext.']
],[.7,1.25,4.7])
heading('Key derivation')
formula('Ka || Kd || Ks || Km = HKDF-SHA256(K, salt="HCADS-v2",\n                                    info="independent-keys", length=120)')
p('HKDF first extracts a pseudorandom key with HMAC, then expands it in numbered blocks. The first 32 bytes become Ka, the next 24 become Kd, and the final two 32-byte slices become Ks and Km. This separation prevents the MAC and encryption operations from directly reusing the same subkey. The implementation is checked against the RFC 5869 test vector.')
heading('Unambiguous encoding')
formula('F(X) = uint32_be(byte_length(X)) || X')
p('Lengths refer to bytes, not JavaScript characters. Each four-byte prefix is big endian. Domain strings and framing distinguish a label from plaintext and distinguish the IV HMAC from the outer MAC. The labels and version are public associated context, not secret data.')
heading('Contexts used by the application')
formula('request:<action>:<challenge>:<field>\nresponse:<action>:<challenge>:<field>\nstore:stdid\nstore:<random row ID>:<field>')
p('Request labels bind the action, one-use challenge and field name. Response labels bind the corresponding action and challenge. The stored student ID uses a fixed context so normalization and deterministic encryption support the UNIQUE constraint. Name, division and branch include the row ID. The database also stores the student ID as a readable username, and passwords use an independent salted hash.')

page('Simple encryption flowchart')
p('Follow the arrows from one plaintext field to an authenticated transport packet. K is either the temporary channel key or the server-only storage key, depending on where the encryption takes place.')
image(encryption,'Figure 1  HC ADS v2 encryption and ciphertext fingerprint',5.85)

page('Every encryption operation explained')
heading('Step 1 Derive four independent keys')
p('Apply the 120-byte HKDF derivation from the preceding section to K. Every subsequent step uses only its designated subkey. HKDF provides key separation; it does not stretch a weak human password into a suitable storage master key. Both master keys in this project are generated from cryptographic randomness or ECDH.')
heading('Step 2 Bind plaintext to its context')
formula('V = HMAC-SHA256(Ks, "HCADS2-SIV" || F(L) || F(P))')
p('This 32-byte value depends on the secret key, field context and complete plaintext. For a fixed K, L and P it is deterministic. A different field label or one-use request challenge changes V. This is SIV-inspired; it is not the standardized AES-SIV construction and must not be called equivalent to it.')
heading('Step 3 Select IVs and apply AES')
formula('IVa = V[0:16]     IVd = V[16:24]\nA = AES-256-CTR(Ka, IVa, P)')
p('AES encrypts counter blocks to make a keystream, which is XORed with P. CTR preserves byte length and needs no padding. The implementation increments the complete 128-bit counter in big-endian order, including carries across words. IV reuse for distinct messages under the same AES key would be dangerous; synthetic IV generation is intended to bind it to the plaintext and context, but this custom construction has not been formally proven.')
heading('Step 4 Pad and apply the legacy cascade')
formula('n = 8 - (byte_length(A) mod 8)\nA_padded = A || byte(n) repeated n times\nB = 3DES-EDE3-CBC(Kd, IVd, A_padded)')
p('PKCS7 adds between one and eight bytes; an aligned value receives a complete padding block. Triple DES encrypts the padded AES output in CBC mode, which uses eight-byte blocks. Padding is applied once and removed strictly once during authenticated decryption. The cascade preserves the requested lab experiment. It does not demonstrate that two ciphers provide a security level equal to the sum of their key sizes.')
heading('Step 5 Authenticate the complete envelope')
formula('body = 0x02 || V || B\nT = HMAC-SHA256(Km, "HCADS2-AUTH" || F(L) || body)\nct = Base64URL(body || T), without padding')
p('The outer tag binds the version, context, synthetic IV and encrypted body using Km. A receiver checks it before invoking CBC or examining padding. Modifying any covered byte, switching field labels or substituting the key causes authentication failure. Base64URL is only a reversible transport encoding; it adds no security.')
heading('Step 6 Compute the fingerprint')
formula('digest = lowercase_hex(SHA256(UTF8(ct)))')
p('The digest hashes the encoded ciphertext text, not the plaintext and not the decoded binary. The client sends {ct, digest}. An attacker can recompute an unkeyed digest, so authenticity depends on T. The server also decrypts, re-encrypts using the same context and key, and compares the new digest to retain the experiment workflow.')

page('Authenticated decryption flowchart')
p('No candidate plaintext is returned until both keyed checks pass. Bad padding, unsupported versions, malformed text and failed tags share one cipher error.')
image(decryption,'Figure 2  Authentication gates before and after decryption',5.65)
p('The byte layout is version 1 byte, V 32 bytes, B a positive multiple of 8 bytes, and T 32 bytes. The minimum envelope is 73 bytes. For plaintext length n, the binary envelope length is 65 + 8 times (floor(n / 8) + 1). Ciphertext length therefore still reveals a padded length interval.')

page('Browser channel and account lifecycle')
heading('Establish a temporary encrypted channel')
p('The browser generates an ephemeral P-256 ECDH key pair using WebCrypto and sends only its public JWK to POST /api/channel. The server generates an ephemeral pair, validates the browser public key, computes the ECDH secret and returns its public JWK plus a random public salt. Both sides derive the same 32-byte channel key using native HKDF-SHA256 with info "AC-channel-v2". The salt is the UTF-8 text of the returned Base64URL token.')
p('A separate 256-bit channel cookie is HttpOnly and SameSite Strict; only its SHA-256 identifier is kept in the server map. The channel key stays in browser closure memory and the server map for ten minutes. It is not written to localStorage, sessionStorage, source files or the database. ECDH alone does not authenticate the server: HTTPS is required outside the loopback demonstration.')
heading('Register and deliver the generated password')
p('1. The browser collects student ID, full name, division and branch. It obtains a 256-bit random challenge for action register. The challenge expires after 60 seconds and belongs to that browser channel.')
p('2. Each field is sealed with Kchannel using request:register:<challenge>:<field>. The server consumes the challenge before cryptographic processing, validates field names, verifies digests and MACs, decrypts and performs the re-encryption digest check.')
p('3. The server normalizes student ID and division to uppercase, trims fields, normalizes the name to NFC and validates their formats. It checks for a duplicate normalized username using a prepared SQL query, then generates the password.')
p('4. The server encrypts profile fields under Kstore, hashes the password with native scrypt and saves username, hash and encrypted profile fields. It sends the password sealed under Kchannel using response:register:<challenge>:pwd. Plaintext passwords are never persisted or logged.')
p('5. The browser authenticates and decrypts the reply and renders a QR locally. The QR contains the plaintext password, so the user must save it privately. Leaving the screen requires a saved-password confirmation; the password and QR are cleared when leaving. A page close or reload loses the one-time display.')
heading('Sign in and access the profile')
p('A fresh login challenge binds encrypted student ID and password to one request. The server authenticates both fields, normalizes the ID and looks up the username. It runs the submitted password through scrypt with the stored salt and parameters, then compares the 32-byte result in constant time. Unknown usernames use a dummy hash. No decrypted password is recovered from the database.')
p('Successful login rotates the current authentication cookie and creates a thirty-minute server session. The encrypted result is bound to the login challenge. Profile and logout requests also require fresh challenges and authenticated proof fields. Profile data is decrypted from storage and re-encrypted for the browser. Logout revokes the session; a restart invalidates all sessions and challenges.')

page('Password generation and quantum randomness')
p('The password has 16 characters selected from a 70-symbol alphabet. At least one uppercase letter, lowercase letter, digit and symbol must appear. The diagram shows the source selection and sampling process.')
image(password_flow,'Figure 3  Password generation with an optional genuine quantum source',4.65)
p('The alphabet is A-Z, a-z, 0-9 and ! @ # $ % & * ?. Only bytes below 210 are accepted because floor(256 / 70) times 70 equals 210. Each symbol then has three equally likely accepted byte values. Bytes 210 through 255 are rejected, removing modulo bias.')
p('Before class filtering, the ideal upper-bound entropy is 16 log2(70), about 98.1 bits; filtering reduces this slightly. HKDF expands a mixed 64-byte seed to a 128-byte candidate stream with salt "HCADS-password-v2" and info "password:<counter>". Missing classes trigger a fresh counter. The generator bounds attempts and clears mutable seed buffers after success.')

page('Database schema and API contract')
heading('Persisted values')
table(['Column','Storage','Purpose'],[
 ['id','Random 128-bit hexadecimal ID','Primary key and row-specific field context.'],['username','Readable normalized student ID, UNIQUE','Login lookup and local administrative display.'],['stdid_ct','Authenticated ciphertext, UNIQUE','Encrypted student ID retained for the lab.'],['name_ct','Authenticated ciphertext','Normalized full name.'],['div_ct','Authenticated ciphertext','Normalized division.'],['branch_ct','Authenticated ciphertext','Validated branch code.'],['password_hash','Salted scrypt verifier','Password verification without reversible storage.'],['created_at','Timestamp metadata','Creation time, not an encrypted personal field.']
],[1.1,2.05,3.5])
p('The username and metadata are readable. Passwords use salted scrypt hashes; other profile fields remain encrypted. Ciphertext digests are transport values and are not stored as password verifiers. Bound SQL parameters and UNIQUE constraints protect queries and duplicates. Protect database files and private backups.')
heading('Endpoints')
table(['Method and route','Input','Result'],[
 ['GET /api/health','None','Public mode and quantum configuration status.'],['POST /api/channel','publicKey JWK','Server publicKey, salt, channel cookie, 600-second lifetime.'],['POST /api/challenge','action','One-use challenge and 60-second lifetime.'],['POST /api/register','challenge and 4 sealed fields','201; encrypted password and source status.'],['POST /api/login','challenge and sealed stdid, pwd','Encrypted result and 1800-second auth cookie.'],['POST /api/profile','challenge and sealed proof','Encrypted profile for the current user.'],['POST /api/logout','challenge and sealed proof','Session revoked and encrypted signed-out result.']
],[1.65,2.0,3.0])
heading('Request envelope')
formula('{ "challenge": "<one-use token>", "fields": {\n    "stdid": { "ct": "<Base64URL>", "digest": "<64 hex>" },\n    "pwd":   { "ct": "<Base64URL>", "digest": "<64 hex>" }\n} }')
p('Login uses stdid and pwd; registration uses stdid, name, div and branch. Profile and logout use proof. Requests are same-origin JSON, limited to 16 KB.')

page('Threat controls and failure behavior')
table(['Threat','Implemented control','Boundary'],[
 ['Ciphertext modification','Independent HMAC checked before decryption; strict padding and version validation.','Requires secret MAC key; no formal proof for the full custom construction.'],['Digest replacement','A recomputed SHA-256 digest still cannot forge the outer MAC.','SHA-256 alone is not authentication.'],['Replay and field substitution','One-use challenge, 60-second expiry and context-bound action and field labels.','State is per process; clustered operation needs a shared atomic store.'],['Cross-site requests','Exact Origin and Host, Fetch Metadata checks, SameSite Strict cookies, JSON-only POST.','XSS executing in this origin can act with the browser channel key.'],['Credential guessing','30 auth-related requests per IP per 15 minutes; 10 failed attempts per account.','Limits reset on restart; account blocking can be abused for denial of service.'],['SQL injection and data exposure','Prepared queries, ciphertext fields, no body or password logs, private file paths never served.','The storage key decrypts profile fields. Passwords remain salted hashes; a compromised server can still intercept submitted credentials.'],['Session theft','256-bit random cookies, HttpOnly, SameSite Strict, Secure under HTTPS; bounded expiry.','These controls do not replace TLS or endpoint security.'],['Script injection','Self-only CSP, local dependencies, textContent for profile rendering.','Browser extensions, device malware and malicious first-party code are outside this protection.']
],[1.15,2.8,2.7])
heading('Uniform cryptographic failure')
p('The shared cipher returns one error for malformed encoding, bad version, invalid lengths, bad tags, padding failure and invalid UTF-8. The server responds with a generic invalid-credentials-or-request error for login failures, rather than naming the wrong credential. Comparisons avoid early byte exits; Node uses timingSafeEqual for ciphertext comparison. JavaScript JIT execution is not guaranteed constant time, and different API paths can still have measurable timing differences.')
heading('Request and state limits')
p('API traffic is limited to 120 requests per IP per minute. Channels are capped at 1000, sessions at 1000, and each channel has at most eight pending challenges. Expired state is purged on a timer and during access. The registration lab caps enrollment at 1000 records. This cap is not a formal cumulative TDEA block-usage policy.')

page('Security limits and source handling')
heading('Legacy primitives and research status')
p('NIST disallowed TDEA for new encryption after 31 December 2023. Keeping Triple DES satisfies this educational experiment, not current approval for production encryption. The browser legacy cipher library CryptoJS 4.2.0 is discontinued. Native WebCrypto and Node crypto are used for ECDH and channel HKDF, but the portable cascade still relies on that legacy library. Passing tests and a clean dependency advisory scan are evidence of implementation behavior, not a cryptographic audit.')
p('CryptoJS project status: https://github.com/brix/crypto-js')
heading('What encryption does not hide')
p('Deterministic encrypted student IDs reveal equality and allow record linkage across snapshots. Row contexts prevent direct equality leakage for other fields across rows, but padded ciphertext lengths remain visible. Registration duplicate responses can reveal that an ID is enrolled. Usernames are intentionally readable. Password hashes cannot be decrypted with Kstore. Legacy migration backups retain the earlier reversible password ciphertext and must remain private. Plaintext necessarily exists in process and browser memory during validation and QR rendering; garbage-collected strings cannot be reliably erased.')
heading('What the project does not implement')
p('There is no MFA, student identity proofing, password recovery, managed key rotation or production multi-server state backend. In-memory sessions disappear on restart. The channel cookie is shared by tabs, so establishing a new channel in one tab can invalidate another. A registration response lost after insertion cannot be recovered from the interface. Retain the one-time credential before leaving.')
heading('Genuine quantum source and fallback')
p('The optional source is the official ANU endpoint https://api.quantumnumbers.anu.edu.au/?length=64&type=uint8 with a server-side x-api-key header. Requests have a 3.5-second timeout and reject redirects. The client validates HTTP success, success=true, exactly 64 integers and values between 0 and 255. The server obtains fresh independent CSPRNG bytes after the quantum request; an invalid external source cannot replace that local source.')
p('Successful integration reports csprng+anu-qrng. A missing key reports csprng-only. An unavailable or malformed response reports csprng-only-fallback. No pseudorandom simulation is described as physical quantum randomness. A genuine live request was not verified without a user-provided API key; the automated tests cover valid, invalid and unavailable mocked responses.')
heading('Production successor')
p('A deployment intended to protect real credentials should replace the custom cascade with reviewed native AEAD, retain strong salted password verification, isolate keys in a managed secret service, enforce trusted HTTPS termination and use a shared atomic session, challenge and throttle store. Add MFA, recovery and operational monitoring as product requirements. The local educational app uses HTTP only on loopback; production configuration rejects an HTTP origin.')

page('Running and validating the project')
heading('Local setup')
p('Use Node.js 24 or newer. From the project folder run the following commands, then open http://localhost:3000. Use the configured hostname exactly; Host and Origin checks are intentional.')
formula('npm install\nnpm run setup\nnpm start')
p('The setup script creates a random storage key only when .env does not already exist. Keep .env private and back it up separately with access controls. Do not regenerate the key for an existing database. PORT and APP_ORIGIN must match. Database files are in data; the report is in docs and is downloadable from the app. npm install builds the local QR asset, so no external browser CDN is required.')
heading('Validation commands and evidence')
formula('npm test\nnpm run check\nnpm audit')
p('The test suite includes 19 automated cases: round trips and Unicode, deterministic context separation, every-bit tampering across the envelope, recomputed-digest attacks, RFC 5869 HKDF, independent native OpenSSL parity, browser module parity, AES counter carry, password quality and quantum responses, plus API lifecycle, salted password verification, legacy migration and portable account export. Tests use temporary databases, not real student records.')
p('The browser workflow is also checked against an isolated in-memory server using an explicit synthetic password. This validates enrollment, QR delivery, saved-password confirmation, sign-in, profile rendering and logout. The fixture script must never be used as the deployed server. Automated tests do not establish resistance to every side channel or attacker, and the live ANU integration requires a real API key.')
heading('Source map')
table(['File','Implementation responsibility'],[
 ['shared/hcads.cjs','Versioned encryption, framing, HKDF, strict authenticated decryption and digest envelopes.'],['server/app.js','Channel exchange, request freshness, API validation, storage translation, sessions and headers.'],['server/password.js','CSPRNG and optional ANU mixing, HKDF candidate generation and unbiased symbol sampling.'],['server/db.js','SQLite schema, username and hash storage, migration and prepared queries.'],['server/password-hash.js','Native salted scrypt generation and verification.'],['public/app.js','Browser channel setup, forms, QR delivery, profile display and safe rendering.']
],[1.7,4.95])
heading('Primary references')
for reference in [
    'RFC 5869  HKDF specification  https://www.rfc-editor.org/rfc/rfc5869',
    'RFC 7366  Encrypt then MAC principle  https://www.rfc-editor.org/rfc/rfc7366',
    'NIST  TDEA withdrawal  https://csrc.nist.gov/News/2023/nist-to-withdraw-sp-800-67-rev-2',
    'ANU  API documentation  https://quantumnumbers.anu.edu.au/documentation'
]:
    para=p(reference);para.paragraph_format.space_after=Pt(3)
    for run in para.runs:run.font.size=Pt(9)

page('Password storage and portable SQLite data')
heading('Salted password verification')
p('The database stores scrypt$32768$8$3$<salt hex>$<hash hex>. Each password receives a random 16-byte salt. Native scrypt uses N=32768, r=8, p=3 and produces 32 bytes. This is the OWASP 32 MiB configuration. SHA-256 remains the transport fingerprint and is not used alone for password storage. Password hashing has at most two active jobs and eight queued jobs.')
p('During login, the backend authenticates and decrypts the request, reads the normalized username, derives the candidate scrypt result with the saved salt and compares it to the saved hash. Correct credentials create a session. Password hashes cannot be decrypted; the generated plaintext is delivered only once through the encrypted response and local QR.')
heading('Display the backend records')
formula('npm run db:view\nnpm run db:view -- --raw')
p('The first command prints usernames, profile fields and hashes locally. The second shows exact database columns. A graphical SQLite viewer can open data/students.sqlite: choose Open Database, Browse Data and the students table. Use read-only mode for demonstrations. No public API exposes the full database or password hashes.')
heading('Run on another laptop')
formula('git clone https://github.com/KrisGod21/SecureAuth.git\ncd SecureAuth\nnpm ci\nnpm run setup\nnpm start')
p('Install Node.js 24 or newer first. A fresh clone creates an empty database with its own key. To preserve existing accounts, run npm run db:export on the original laptop. It creates a consistent SQLite snapshot and the matching private .env in an ignored private-backups folder. Copy both into a fresh destination installation before setup, then sign in with the existing password. Do not upload this private export to GitHub.')
p('The export uses the SQLite backup API to capture committed WAL data safely. Active sessions are not copied. Two local installations do not automatically synchronize. Startup migrates an older pwd_ct database in a transaction and creates a private legacy backup first. The current database removes the reversible password column; old backups may still contain it.')
p('Full instructions are in docs/DATABASE_AND_TRANSFER.md. Hosting still requires a persistent disk or an external database and shared state; this local SQLite version is not a complete Vercel serverless deployment.')
p('OWASP scrypt configuration: https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html#scrypt')
path=ROOT/'docs'/'AC_Algorithm_and_Application.docx'
doc.save(path)
print(f'Created {path}')
